"""Spreadsheet import: turn an uploaded .xlsx into rows for the column-mapping UI."""

import datetime
import zipfile

from flask import Blueprint, jsonify, request
from openpyxl import load_workbook

from common import json_error

bp = Blueprint("imports", __name__, url_prefix="/api")

MAX_XLSX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.isoformat()
    return str(value).strip()


@bp.route("/import/xlsx", methods=["POST"])
def import_xlsx():
    """
    Returns {"rows": [[header...], [cell...], ...]} from the first worksheet,
    the same shape the frontend's CSV parser produces.
    """
    file = request.files.get("file")
    if not file or not file.filename:
        return json_error("No file provided")
    if not file.filename.lower().endswith(".xlsx"):
        return json_error("Only .xlsx files are supported (save older .xls files as .xlsx or CSV)")

    file.seek(0, 2)
    if file.tell() > MAX_XLSX_BYTES:
        return json_error("Spreadsheet is larger than 5MB")
    file.seek(0)

    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError):
        return json_error("Couldn't read that file — is it a valid .xlsx spreadsheet?")

    try:
        sheet = workbook.worksheets[0]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell_text(v) for v in row]
            if any(cells):
                rows.append(cells)
            if len(rows) > MAX_ROWS:
                return json_error(f"Spreadsheet has more than {MAX_ROWS} rows")
    finally:
        workbook.close()

    # Pad ragged rows so every row has as many cells as the widest one
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return jsonify({"rows": rows, "sheet": sheet.title})
